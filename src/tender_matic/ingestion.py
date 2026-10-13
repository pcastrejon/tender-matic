"""Local document intake for phase 1; this module never calls a model or network."""

from __future__ import annotations

import hashlib
import json
import os
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterator

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


@dataclass(frozen=True)
class ExtractionWarning:
    code: str
    message: str
    source_location: str | None = None


@dataclass(frozen=True)
class NormalizedPassage:
    passage_id: str
    text: str
    source_location: str


@dataclass
class NormalizedDocument:
    document_id: str
    relative_path: str
    format: str
    status: str
    passages: list[NormalizedPassage] = field(default_factory=list)
    warnings: list[ExtractionWarning] = field(default_factory=list)


@dataclass
class FileManifestEntry:
    relative_path: str
    format: str
    status: str
    size_bytes: int | None
    document_id: str | None
    passage_count: int
    warnings: list[ExtractionWarning] = field(default_factory=list)


@dataclass
class IngestionResult:
    input_root: str
    manifest: list[FileManifestEntry]
    documents: list[NormalizedDocument]
    warnings: list[ExtractionWarning] = field(default_factory=list)

    def to_manifest_dict(self) -> dict[str, Any]:
        return {
            "schema_version": "1",
            "input_root": self.input_root,
            "files": [_serialize(entry) for entry in self.manifest],
            "warnings": [_serialize(warning) for warning in self.warnings],
        }


class UnsupportedEncryptedPdfError(ValueError):
    """Raised when a PDF cannot be parsed without a password."""


def _serialize(value: Any) -> dict[str, Any]:
    return asdict(value)


def _normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _passage(document_id: str, text: str, location: str) -> NormalizedPassage | None:
    normalized = _normalize_text(text)
    if not normalized:
        return None
    passage_id = hashlib.sha256(
        f"{document_id}\0{location}\0{normalized}".encode("utf-8")
    ).hexdigest()[:20]
    return NormalizedPassage(passage_id, normalized, location)


def _document_id(relative_path: str, raw_bytes: bytes) -> str:
    identity = relative_path.encode("utf-8") + b"\0" + raw_bytes
    return hashlib.sha256(identity).hexdigest()[:20]


def _text_passages(
    raw_bytes: bytes, document_id: str
) -> tuple[list[NormalizedPassage], list[ExtractionWarning]]:
    warnings: list[ExtractionWarning] = []
    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw_bytes.decode("utf-8-sig", errors="replace")
        warnings.append(
            ExtractionWarning(
                "invalid_utf8_replaced",
                "Invalid UTF-8 bytes were replaced during text decoding.",
            )
        )

    passages = [
        parsed
        for line_number, line in enumerate(text.splitlines(), start=1)
        if (parsed := _passage(document_id, line, f"line:{line_number}"))
        is not None
    ]
    if not passages:
        warnings.append(ExtractionWarning("no_text_extracted", "No text was extracted."))
    return passages, warnings


def _pdf_passages(
    path: Path, document_id: str
) -> tuple[list[NormalizedPassage], list[ExtractionWarning]]:
    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        raise UnsupportedEncryptedPdfError

    passages: list[NormalizedPassage] = []
    warnings: list[ExtractionWarning] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            warnings.append(
                ExtractionWarning(
                    "page_extraction_failed",
                    "Text extraction failed for this PDF page.",
                    f"page:{page_number}",
                )
            )
            continue

        parsed = _passage(document_id, text, f"page:{page_number}")
        if parsed is None:
            warnings.append(
                ExtractionWarning(
                    "ocr_required_or_empty_page",
                    "No text layer was found on this page; OCR is not included.",
                    f"page:{page_number}",
                )
            )
        else:
            passages.append(parsed)

    if not passages and not warnings:
        warnings.append(ExtractionWarning("no_text_extracted", "No text was extracted."))
    return passages, warnings


def _iter_docx_container(container: Any, prefix: str) -> Iterator[tuple[str, str]]:
    paragraph_number = 0
    table_number = 0
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_number += 1
            yield f"{prefix}paragraph:{paragraph_number}", block.text
        elif isinstance(block, Table):
            table_number += 1
            for row_number, row in enumerate(block.rows, start=1):
                for cell_number, cell in enumerate(row.cells, start=1):
                    yield (
                        f"{prefix}table:{table_number}:row:{row_number}:cell:{cell_number}",
                        " ".join(paragraph.text for paragraph in cell.paragraphs),
                    )


def _iter_docx_blocks(document: Any) -> Iterator[tuple[str, str]]:
    yield from _iter_docx_container(document, "")
    for section_number, section in enumerate(document.sections, start=1):
        for story_name in ("header", "first_page_header", "even_page_header", "footer",
                           "first_page_footer", "even_page_footer"):
            story = getattr(section, story_name)
            if story.is_linked_to_previous and section_number > 1:
                continue
            yield from _iter_docx_container(story, f"{story_name}:section:{section_number}:")


def _docx_passages(
    path: Path, document_id: str
) -> tuple[list[NormalizedPassage], list[ExtractionWarning]]:
    document = Document(path)
    passages = [
        parsed
        for location, text in _iter_docx_blocks(document)
        if (parsed := _passage(document_id, text, location)) is not None
    ]
    warnings: list[ExtractionWarning] = []
    if document.inline_shapes:
        warnings.append(
            ExtractionWarning(
                "embedded_objects_not_parsed",
                "Embedded images or objects were not inspected.",
            )
        )
    if not passages:
        warnings.append(ExtractionWarning("no_text_extracted", "No text was extracted."))
    return passages, warnings


def _status(passage_count: int, warnings: list[ExtractionWarning]) -> str:
    if not passage_count:
        return "partial"
    return "partial" if warnings else "parsed"


def _failure_entry(relative_path: str, file_format: str, code: str, message: str,
                   size_bytes: int | None = None) -> FileManifestEntry:
    return FileManifestEntry(
        relative_path=relative_path,
        format=file_format,
        status="failed",
        size_bytes=size_bytes,
        document_id=None,
        passage_count=0,
        warnings=[ExtractionWarning(code, message)],
    )


def ingest_directory(input_root: Path) -> IngestionResult:
    """Parse supported files below input_root without modifying source files."""
    resolved_root = input_root.resolve(strict=True)
    if not resolved_root.is_dir():
        raise NotADirectoryError(f"Input path is not a directory: {input_root}")

    manifest: list[FileManifestEntry] = []
    documents: list[NormalizedDocument] = []
    run_warnings: list[ExtractionWarning] = []
    discovered_files = sorted(resolved_root.rglob("*"), key=lambda path: path.as_posix())

    for path in discovered_files:
        if not path.is_file():
            if not path.is_symlink():
                continue
        relative_path = path.relative_to(resolved_root).as_posix()
        file_format = path.suffix.casefold().lstrip(".") or "unknown"

        if path.is_symlink():
            manifest.append(
                _failure_entry(
                    relative_path,
                    file_format,
                    "symlink_not_allowed",
                    "Symbolic links are not ingested.",
                )
            )
            continue

        if path.suffix.casefold() not in SUPPORTED_EXTENSIONS:
            manifest.append(
                FileManifestEntry(
                    relative_path=relative_path,
                    format=file_format,
                    status="unsupported",
                    size_bytes=path.stat().st_size,
                    document_id=None,
                    passage_count=0,
                    warnings=[
                        ExtractionWarning(
                            "unsupported_format",
                            "Supported formats are TXT, Markdown, PDF, and DOCX.",
                        )
                    ],
                )
            )
            continue

        try:
            raw_bytes = path.read_bytes()
            size_bytes = len(raw_bytes)
            document_id = _document_id(relative_path, raw_bytes)
            suffix = path.suffix.casefold()
            if suffix in {".txt", ".md"}:
                passages, warnings = _text_passages(raw_bytes, document_id)
            elif suffix == ".pdf":
                passages, warnings = _pdf_passages(path, document_id)
            else:
                passages, warnings = _docx_passages(path, document_id)
        except Exception as error:
            message = (
                "Encrypted PDF files are not supported."
                if isinstance(error, UnsupportedEncryptedPdfError)
                else "The file could not be parsed. Verify that it is valid and readable."
            )
            manifest.append(
                _failure_entry(
                    relative_path,
                    file_format,
                    "parse_failed",
                    message,
                    path.stat().st_size if path.exists() else None,
                )
            )
            continue

        status = _status(len(passages), warnings)
        parsed_document = NormalizedDocument(
            document_id=document_id,
            relative_path=relative_path,
            format=file_format,
            status=status,
            passages=passages,
            warnings=warnings,
        )
        documents.append(parsed_document)
        manifest.append(
            FileManifestEntry(
                relative_path=relative_path,
                format=file_format,
                status=status,
                size_bytes=size_bytes,
                document_id=document_id,
                passage_count=len(passages),
                warnings=warnings,
            )
        )

    if not discovered_files:
        run_warnings.append(
            ExtractionWarning("empty_input_directory", "No files were found in the input directory.")
        )

    return IngestionResult(
        input_root=str(resolved_root),
        manifest=manifest,
        documents=documents,
        warnings=run_warnings,
    )


def write_artifacts(result: IngestionResult, output_root: Path) -> None:
    """Write normalized JSON artifacts outside the source input tree."""
    input_root = Path(result.input_root).resolve()
    resolved_output = output_root.resolve()
    if resolved_output == input_root or input_root in resolved_output.parents:
        raise ValueError("Output directory must be outside the input directory")
    if resolved_output in input_root.parents:
        raise ValueError("Output directory must not contain the input directory")

    documents_directory = resolved_output / "documents"
    documents_directory.mkdir(parents=True, exist_ok=True)
    manifest_path = resolved_output / "manifest.json"
    _write_json(manifest_path, result.to_manifest_dict())
    for document in result.documents:
        _write_json(documents_directory / f"{document.document_id}.json", _serialize(document))


def _write_json(path: Path, value: dict[str, Any]) -> None:
    temporary_path = path.with_suffix(path.suffix + ".tmp")
    temporary_path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + os.linesep,
        encoding="utf-8",
    )
    temporary_path.replace(path)