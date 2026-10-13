import json
import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from tender_matic.ingest import main as ingestion_cli_main
from tender_matic.ingestion import ingest_directory, write_artifacts


def create_text_pdf(path: Path, pages: list[str]) -> None:
    pdf_canvas = canvas.Canvas(str(path))
    for text in pages:
        pdf_canvas.drawString(72, 720, text)
        pdf_canvas.showPage()
    pdf_canvas.save()


class TextDocumentIngestionTests(unittest.TestCase):
    def test_parses_text_and_markdown_with_line_references(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "notes.txt").write_text(" First line  \n\nSecond\tline\n", encoding="utf-8")
            (root / "requirements.md").write_text("# Scope\n\nMust support export.\n", encoding="utf-8")

            result = ingest_directory(root)

            self.assertEqual([entry.status for entry in result.manifest], ["parsed", "parsed"])
            first_document = result.documents[0]
            self.assertEqual(first_document.passages[0].text, "First line")
            self.assertEqual(first_document.passages[0].source_location, "line:1")
            self.assertEqual(first_document.passages[1].source_location, "line:3")
            markdown_document = result.documents[1]
            self.assertEqual(markdown_document.passages[0].source_location, "line:1")

    def test_invalid_utf8_is_replaced_and_warned(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "invalid.txt").write_bytes(b"valid \xff text")

            result = ingest_directory(root)

            self.assertEqual(result.manifest[0].status, "partial")
            self.assertIn("invalid_utf8_replaced", {warning.code for warning in result.manifest[0].warnings})


class PdfIngestionTests(unittest.TestCase):
    def test_parses_text_pdf_with_page_references(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            create_text_pdf(root / "tender.pdf", ["Requirement A", "Requirement B"])

            result = ingest_directory(root)

            document = result.documents[0]
            self.assertEqual(document.status, "parsed")
            self.assertEqual([passage.source_location for passage in document.passages], ["page:1", "page:2"])
            self.assertIn("Requirement A", document.passages[0].text)

    def test_scanned_or_empty_pdf_page_is_flagged_for_ocr(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            create_text_pdf(root / "scan.pdf", [""])

            result = ingest_directory(root)

            self.assertEqual(result.manifest[0].status, "partial")
            self.assertIn(
                "ocr_required_or_empty_page",
                {warning.code for warning in result.manifest[0].warnings},
            )

    def test_encrypted_pdf_is_reported_as_failed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            plain_pdf = root / "plain.pdf"
            create_text_pdf(plain_pdf, ["Synthetic content"])
            reader = PdfReader(plain_pdf)
            writer = PdfWriter()
            writer.append_pages_from_reader(reader)
            writer.encrypt("synthetic-password")
            encrypted_pdf = root / "encrypted.pdf"
            with encrypted_pdf.open("wb") as output:
                writer.write(output)
            plain_pdf.unlink()

            result = ingest_directory(root)

            self.assertEqual(result.manifest[0].status, "failed")
            self.assertIn("Encrypted PDF", result.manifest[0].warnings[0].message)


class DocxIngestionTests(unittest.TestCase):
    def test_parses_paragraphs_and_table_cells_with_locations(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "spec.docx"
            document = Document()
            document.add_paragraph("Tender scope")
            document.sections[0].header.paragraphs[0].text = "Synthetic header"
            document.sections[0].footer.paragraphs[0].text = "Synthetic footer"
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Requirement"
            table.cell(0, 1).text = "Must export"
            document.save(path)

            result = ingest_directory(root)

            parsed = result.documents[0]
            self.assertEqual(parsed.status, "parsed")
            self.assertEqual(parsed.passages[0].source_location, "paragraph:1")
            self.assertEqual(parsed.passages[1].source_location, "table:1:row:1:cell:1")
            self.assertEqual(parsed.passages[2].text, "Must export")
            locations = {passage.source_location for passage in parsed.passages}
            self.assertIn("header:section:1:paragraph:1", locations)
            self.assertIn("footer:section:1:paragraph:1", locations)


class ManifestAndOutputTests(unittest.TestCase):
    def test_cli_creates_manifest_and_document_artifact(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root = base / "input"
            output_root = base / "output"
            input_root.mkdir()
            (input_root / "requirements.txt").write_text("Synthetic requirement", encoding="utf-8")
            arguments = [
                "tender-matic-ingest",
                "--input",
                str(input_root),
                "--output",
                str(output_root),
            ]

            with patch("sys.argv", arguments), redirect_stdout(io.StringIO()):
                exit_code = ingestion_cli_main()

            self.assertEqual(exit_code, 0)
            self.assertTrue((output_root / "manifest.json").is_file())
            self.assertEqual(len(list((output_root / "documents").glob("*.json"))), 1)

    def test_unsupported_and_corrupt_files_are_explicit(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "old.doc").write_bytes(b"not parsed")
            (root / "broken.pdf").write_bytes(b"not a pdf")

            with self.assertLogs("pypdf", level="WARNING"):
                result = ingest_directory(root)

            self.assertEqual([entry.status for entry in result.manifest], ["failed", "unsupported"])
            self.assertEqual(result.manifest[0].warnings[0].code, "parse_failed")
            self.assertEqual(result.manifest[1].warnings[0].code, "unsupported_format")

    def test_writes_manifest_and_normalized_documents_outside_input(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root = base / "input"
            output_root = base / "artifacts"
            input_root.mkdir()
            (input_root / "tender.txt").write_text("Synthetic requirement", encoding="utf-8")

            result = ingest_directory(input_root)
            write_artifacts(result, output_root)

            manifest = json.loads((output_root / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest["schema_version"], "1")
            document_files = list((output_root / "documents").glob("*.json"))
            self.assertEqual(len(document_files), 1)
            document = json.loads(document_files[0].read_text(encoding="utf-8"))
            self.assertEqual(document["passages"][0]["text"], "Synthetic requirement")

    def test_rejects_output_directory_inside_source_tree(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            input_root = Path(directory) / "input"
            input_root.mkdir()
            result = ingest_directory(input_root)

            with self.assertRaises(ValueError):
                write_artifacts(result, input_root / "artifacts")

    def test_symlink_is_reported_and_not_followed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            input_root = base / "input"
            input_root.mkdir()
            outside = base / "outside.txt"
            outside.write_text("must not be read", encoding="utf-8")
            (input_root / "linked.txt").symlink_to(outside)

            result = ingest_directory(input_root)

            self.assertEqual(result.manifest[0].status, "failed")
            self.assertEqual(result.manifest[0].warnings[0].code, "symlink_not_allowed")
            self.assertEqual(result.documents, [])


if __name__ == "__main__":
    unittest.main()